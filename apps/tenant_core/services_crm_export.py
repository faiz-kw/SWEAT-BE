"""
apps/tenant_core/services_crm_export.py — Dynamic Excel Export Service for CRM Leads & Voice Calls.

Exports leads, complete CRM profile, Meta Lead Form question/answers,
and Sarvam AI Voice Calling session links (audio recording download URLs,
AI call summaries, dialogue transcripts) into a multi-tab formatted Excel (.xlsx) file.
"""

import io
import os
import requests
from datetime import datetime, timezone, timedelta
from typing import Optional, Dict, Any

from django.conf import settings
from django.utils import timezone as dj_timezone

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from apps.tenant_core.models_crm import (
    Lead, CallSession, TrialBooking, SalesFollowupTask, LeadNote
)
from apps.tenant_core.models_meta_leads import MetaLeadImport
from apps.tenant_core.models_infra import Integration

# Indian Standard Time (IST)
IST = timezone(timedelta(hours=5, minutes=30))


def fmt_dt(dt):
    """Format datetime cleanly in IST."""
    if not dt:
        return ""
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(IST).strftime("%Y-%m-%d %I:%M:%S %p IST")


def fetch_sarvam_interactions(cfg: Dict[str, Any], api_key: str, start_time: datetime) -> Dict[str, Any]:
    """Fetch recent Sarvam call interactions from Sarvam Analytics API."""
    org_id = cfg.get('org_id') or getattr(settings, 'SARVAM_ORG_ID', '') or os.getenv('SARVAM_ORG_ID', '')
    workspace_id = cfg.get('workspace_id') or getattr(settings, 'SARVAM_WORKSPACE_ID', '') or os.getenv('SARVAM_WORKSPACE_ID', '')
    app_id = cfg.get('app_id') or getattr(settings, 'SARVAM_APP_ID', '') or os.getenv('SARVAM_APP_ID', '')

    if not (org_id and workspace_id and app_id and api_key):
        return {}

    url = f"https://apps.sarvam.ai/api/analytics/v1/{org_id}/{workspace_id}/{app_id}/interactions"
    params = {
        'start_datetime': (start_time - timedelta(days=2)).strftime('%Y-%m-%dT00:00:00Z'),
        'end_datetime': (datetime.now(timezone.utc) + timedelta(hours=4)).strftime('%Y-%m-%dT%H:%M:%SZ'),
        'include_transcript': 'true',
        'limit': 100,
    }

    by_job = {}
    try:
        resp = requests.get(url, headers={'X-API-Key': api_key}, params=params, timeout=12)
        if resp.status_code == 200:
            items = resp.json().get('items', [])
            for it in items:
                job_id = it.get('job_id')
                if job_id:
                    by_job[job_id] = it
                iid = it.get('interaction_id')
                if iid:
                    by_job[iid] = it
    except Exception:
        pass

    return by_job


def export_leads_to_excel_buffer(leads_qs, db_alias: str, request=None) -> bytes:
    """
    Generates an in-memory .xlsx workbook for the given leads queryset.
    Returns the binary content as bytes.
    """
    # 1. Resolve Sarvam telephony configuration & credentials
    s_integ = Integration.objects.using(db_alias).filter(provider__iexact='SARVAM').first()
    cfg = s_integ.configuration if s_integ else {}
    api_key = (
        getattr(settings, 'SARVAM_API_KEY', '') 
        or os.getenv('SARVAM_API_KEY', '') 
        or 'sk_samvaad_3eda4iiv_dNdBuzj4OTnJC9Vv4G7KBtBS'
    )

    backend_public_url = (
        getattr(settings, 'BACKEND_PUBLIC_URL', '')
        or os.getenv('BACKEND_PUBLIC_URL', '')
        or (request.build_absolute_uri('/') if request else 'http://localhost:8000')
    ).rstrip('/')

    rec_cache_dir = os.path.join(settings.BASE_DIR, 'media', 'recordings')

    # Fetch Sarvam telephony interactions for audio link matching
    earliest_lead = leads_qs.order_by('created_at').first()
    start_utc = earliest_lead.created_at if earliest_lead else (datetime.now(timezone.utc) - timedelta(days=7))
    sarvam_data = fetch_sarvam_interactions(cfg, api_key, start_utc)

    # Pre-fetch optimized lead records
    leads = leads_qs.using(db_alias).select_related(
        'organization', 'branch', 'lead_source', 'interested_program',
        'assigned_sales_user', 'assigned_trainer_user', 'referred_by_user'
    ).prefetch_related(
        'call_sessions', 'notes', 'activities', 'trial_bookings'
    ).order_by('-created_at')

    # Initialize Workbook
    wb = openpyxl.Workbook()
    ws_main = wb.active
    ws_main.title = "CRM Leads Master"

    # Styling
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid") # Navy
    fill_call_header = PatternFill(start_color="0D5C3A", end_color="0D5C3A", fill_type="solid") # Dark Green
    fill_meta_header = PatternFill(start_color="5A3D8A", end_color="5A3D8A", fill_type="solid") # Purple

    font_data = Font(name="Calibri", size=10)
    font_link = Font(name="Calibri", size=10, color="0000FF", underline="single")
    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")
    align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    headers = [
        # 1. Core Lead Identification
        "Lead ID",
        "Created Date (IST)",
        "Updated Date (IST)",
        "Full Name",
        "First Name",
        "Last Name",
        "Phone Number",
        "Email Address",
        "Current CRM Status",
        "Studio / Branch",
        "Organization",

        # 2. Sarvam AI Voice Calling & Audio Links
        "Sarvam Call Status",
        "Connected with Sarvam?",
        "Call Duration (Seconds)",
        "Audio Download Link (Sarvam Cloud)",
        "Audio Stream Link (Backend Web)",
        "Audio File Cached on Server?",
        "Call Started At (IST)",
        "Sarvam Attempt ID",
        "Sarvam Interaction ID",
        "Sarvam AI Call Summary",
        "Call Disposition",
        "Extracted Next Step",
        "Scheduled Next Step Date/Time",
        "Full Call Transcript",

        # 3. Meta Lead Form Answers (Facebook / Instagram)
        "Meta Campaign Name",
        "Meta Ad Name",
        "Preferred Studio (From Form)",
        "Workout Type / Goal (From Form)",
        "Start Timeline (From Form)",
        "Monthly Budget / Investment (From Form)",
        "Preferred Schedule / Slot (From Form)",
        "All Meta Form Answers (Combined)",
        "Meta Form ID",
        "Meta External Lead ID",

        # 4. CRM Commercial & Profile Data
        "Lead Source",
        "First Touch Source",
        "Latest Touch Source",
        "Campaign Reference",
        "Fitness Goal",
        "Interested Program",
        "Gender",
        "Date of Birth",
        "Occupation",
        "Company Name",
        "Locality / Area",
        "Country",
        "Assigned Sales Rep",
        "Assigned Trainer",
        "Referred By",
        "Do Not Contact (DNC)",
        "Consent WhatsApp",
        "Consent Email",
        "Consent SMS",
        "First Response Time (Sec)",
        "Response SLA Status",

        # 5. Trial Bookings & Follow-ups
        "Trial Booked?",
        "Trial Date & Time (IST)",
        "Trial Status",
        "Trial Studio / Branch",
        "Follow-up Task Scheduled?",
        "Follow-up Task Type",
        "Follow-up Due Date (IST)",
        "Follow-up Task Status",
        "CRM Notes / History",
    ]

    ws_main.append(headers)
    ws_main.row_dimensions[1].height = 28
    for col_idx, cell in enumerate(ws_main[1], start=1):
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        col_name = headers[col_idx - 1]
        if "Sarvam" in col_name or "Audio" in col_name or "Call" in col_name:
            cell.fill = fill_call_header
        elif "Meta" in col_name or "(From Form)" in col_name:
            cell.fill = fill_meta_header
        else:
            cell.fill = fill_header

    for r_idx, lead in enumerate(leads, start=2):
        lead_id = str(lead.id)
        full_name = f"{lead.first_name} {lead.last_name}".strip()
        phone = lead.phone_normalized or ""
        email = lead.email_normalized or ""
        branch_name = lead.branch.name if lead.branch else "Organization-Wide"
        org_name = lead.organization.name if lead.organization else "SWEAT"
        lead_source = lead.lead_source.name if lead.lead_source else ""
        assigned_sales = lead.assigned_sales_user.email if lead.assigned_sales_user else ""
        assigned_trainer = lead.assigned_trainer_user.email if lead.assigned_trainer_user else ""
        program_name = lead.interested_program.name if lead.interested_program else ""

        # Meta Form Data
        meta_import = MetaLeadImport.objects.using(db_alias).filter(lead=lead).order_by('-received_at').first()
        if not meta_import and phone:
            meta_import = MetaLeadImport.objects.using(db_alias).filter(field_data__icontains=phone[-10:]).order_by('-received_at').first()

        meta_campaign = ""
        meta_ad = ""
        meta_form_id = ""
        meta_ext_lead_id = ""
        pref_studio = ""
        workout_type = ""
        start_timeline = ""
        monthly_budget = ""
        preferred_slot = ""
        all_answers_lines = []

        if meta_import:
            meta_campaign = meta_import.campaign_name or ""
            meta_ad = meta_import.ad_name or ""
            meta_form_id = meta_import.form_id or ""
            meta_ext_lead_id = meta_import.external_lead_id or ""

            for field in (meta_import.field_data or []):
                name = str(field.get('name', '')).strip()
                vals = field.get('values', [])
                val_str = ", ".join(vals) if vals else ""
                clean_name = name.replace('_', ' ').replace('?', '').strip().title()
                all_answers_lines.append(f"{clean_name}: {val_str}")

                name_l = name.lower()
                if 'studio' in name_l:
                    pref_studio = val_str
                elif 'type' in name_l or 'pilates' in name_l or 'looking' in name_l:
                    workout_type = val_str
                elif 'soon' in name_l or 'plan' in name_l or 'start' in name_l:
                    start_timeline = val_str
                elif 'investment' in name_l or 'budget' in name_l or 'comfortable' in name_l:
                    monthly_budget = val_str
                elif 'week' in name_l or 'timing' in name_l or 'slot' in name_l or 'time' in name_l:
                    preferred_slot = val_str

        all_answers_str = "\n".join(all_answers_lines)

        # Call Sessions & Sarvam Details
        call = lead.call_sessions.filter(provider__iexact='SARVAM').order_by('-created_at').first()
        if not call:
            call = lead.call_sessions.order_by('-created_at').first()

        call_status = ""
        is_connected = "No"
        call_duration = ""
        direct_audio_url = ""
        backend_audio_url = ""
        is_cached_str = "No"
        call_started_at = ""
        attempt_id = ""
        interaction_id = ""
        call_summary = ""
        call_disposition = ""
        next_step = ""
        next_step_dt = ""
        transcript_text = ""

        if call:
            call_status = call.status
            call_duration = round(call.duration_seconds, 1) if call.duration_seconds else 0.0
            attempt_id = call.provider_attempt_id or ""
            call_started_at = fmt_dt(call.started_at or call.created_at)

            if call.status in ('COMPLETED', 'CONNECTED') or (call.duration_seconds and call.duration_seconds > 0):
                is_connected = "Yes"

            local_rec_path = os.path.join(rec_cache_dir, f"{call.id}.wav")
            if os.path.exists(local_rec_path) and os.path.getsize(local_rec_path) > 1000:
                is_cached_str = f"Yes ({round(os.path.getsize(local_rec_path)/1024/1024, 2)} MB)"

            matched_sarvam = sarvam_data.get(call.provider_attempt_id)
            if not matched_sarvam and call.metadata and isinstance(call.metadata, dict):
                iid = call.metadata.get('interaction_id')
                if iid and iid in sarvam_data:
                    matched_sarvam = sarvam_data[iid]

            if matched_sarvam:
                direct_audio_url = matched_sarvam.get('audio_url') or ""
                interaction_id = matched_sarvam.get('interaction_id') or ""
                if not call.transcript and matched_sarvam.get('translate_transcript'):
                    call.transcript = matched_sarvam.get('translate_transcript')

            if not interaction_id and isinstance(call.metadata, dict):
                interaction_id = call.metadata.get('interaction_id') or ""

            if not direct_audio_url and interaction_id and cfg.get('org_id'):
                encoded_iid = str(interaction_id).replace('/', '%2F')
                direct_audio_url = f"https://apps.sarvam.ai/api/analytics/v1/{cfg.get('org_id')}/{cfg.get('workspace_id')}/{cfg.get('app_id')}/recordings/{encoded_iid}"

            if is_connected == "Yes" or call_status == 'COMPLETED':
                backend_audio_url = f"{backend_public_url}/api/v1/voice-agent/recordings/{call.id}/audio/"

            vars_dict = call.final_agent_variables or {}
            call_summary = vars_dict.get('call_summary') or ""
            call_disposition = vars_dict.get('disposition') or ""
            next_step = vars_dict.get('chosen_next_step') or ""
            next_step_dt = vars_dict.get('next_step_date_time') or ""

            if call.transcript and isinstance(call.transcript, list):
                turns = []
                for t in call.transcript:
                    role = "SWEAT Agent (Mahek)" if t.get('role') == 'agent' else "Prospect"
                    text = t.get('text') or t.get('indic_text') or t.get('en_text') or ""
                    turns.append(f"{role}: {text}")
                transcript_text = "\n".join(turns)

        # Trial Bookings
        trial = TrialBooking.objects.using(db_alias).filter(lead=lead).order_by('-created_at').first()
        has_trial = "Yes" if trial else "No"
        trial_dt = fmt_dt(trial.scheduled_start) if trial and trial.scheduled_start else ""
        trial_status = trial.status if trial else ""
        trial_branch = trial.branch.name if trial and trial.branch else ""

        # Follow-up Tasks
        task = SalesFollowupTask.objects.using(db_alias).filter(lead=lead).order_by('-due_at').first()
        has_task = "Yes" if task else "No"
        task_type = task.task_type if task else ""
        task_due = fmt_dt(task.due_at) if task and task.due_at else ""
        task_status = task.status if task else ""

        # Notes
        notes_qs = LeadNote.objects.using(db_alias).filter(lead=lead).order_by('created_at')
        notes_list = []
        for n in notes_qs:
            author = n.created_by_user.email if n.created_by_user else "System"
            notes_list.append(f"[{fmt_dt(n.created_at)} - {n.note_type} - {author}]: {n.note_text}")
        notes_combined = "\n".join(notes_list)

        row = [
            lead_id,
            fmt_dt(lead.created_at),
            fmt_dt(lead.updated_at),
            full_name,
            lead.first_name,
            lead.last_name,
            phone,
            email,
            lead.current_status,
            branch_name,
            org_name,

            call_status,
            is_connected,
            call_duration,
            direct_audio_url if direct_audio_url else ("N/A (No audio/Not connected)" if is_connected == "No" else "Processing"),
            backend_audio_url if backend_audio_url else "N/A",
            is_cached_str,
            call_started_at,
            attempt_id,
            interaction_id,
            call_summary,
            call_disposition,
            next_step,
            next_step_dt,
            transcript_text,

            meta_campaign,
            meta_ad,
            pref_studio,
            workout_type,
            start_timeline,
            monthly_budget,
            preferred_slot,
            all_answers_str,
            meta_form_id,
            meta_ext_lead_id,

            lead_source,
            lead.first_touch_source or "",
            lead.latest_touch_source or "",
            lead.campaign_reference or "",
            lead.fitness_goal or "",
            program_name,
            lead.gender or "",
            str(lead.date_of_birth) if lead.date_of_birth else "",
            lead.occupation or "",
            lead.company_name or "",
            lead.area or "",
            lead.country or "",
            assigned_sales,
            assigned_trainer,
            lead.referred_by_name or "",
            "Yes" if lead.do_not_contact else "No",
            "Yes" if lead.consent_whatsapp else "No",
            "Yes" if lead.consent_email else "No",
            "Yes" if lead.consent_sms else "No",
            lead.first_response_time_seconds or "",
            lead.response_sla_status,

            has_trial,
            trial_dt,
            trial_status,
            trial_branch,
            has_task,
            task_type,
            task_due,
            task_status,
            notes_combined,
        ]

        ws_main.append(row)
        ws_main.row_dimensions[r_idx].height = 22

        for col_i in range(1, len(row) + 1):
            cell = ws_main.cell(row=r_idx, column=col_i)
            cell.font = font_data
            cell.border = thin_border
            cell.alignment = align_left

        if direct_audio_url and direct_audio_url.startswith("http"):
            cell_direct = ws_main.cell(row=r_idx, column=15)
            cell_direct.value = "⬇ Download Sarvam Audio"
            cell_direct.hyperlink = direct_audio_url
            cell_direct.font = font_link

        if backend_audio_url and backend_audio_url.startswith("http"):
            cell_backend = ws_main.cell(row=r_idx, column=16)
            cell_backend.value = "▶ Play / Stream Audio"
            cell_backend.hyperlink = backend_audio_url
            cell_backend.font = font_link

        ws_main.cell(row=r_idx, column=9).alignment = align_center
        cell_conn = ws_main.cell(row=r_idx, column=13)
        cell_conn.alignment = align_center
        if is_connected == "Yes":
            cell_conn.fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")

    for col in ws_main.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if '\n' in val:
                val = val.split('\n')[0]
            if len(val) > max_len:
                max_len = len(val)
        ws_main.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)

    # Sheet 2: Sarvam Voice Calls
    ws_calls = wb.create_sheet(title="Sarvam Voice Calls")
    call_headers = [
        "Lead Name", "Phone Number", "Call Status", "Duration (s)",
        "Direct Audio Download Link", "Backend Audio URL", "Cached On Disk?",
        "Attempt ID", "Interaction ID", "Call Started At (IST)",
        "AI Call Summary", "Disposition", "Next Step", "Next Step Date/Time",
        "Full Dialogue Transcript"
    ]
    ws_calls.append(call_headers)
    ws_calls.row_dimensions[1].height = 26
    for c in ws_calls[1]:
        c.font = font_header
        c.fill = fill_call_header
        c.alignment = Alignment(horizontal="center", vertical="center")

    call_row_idx = 2
    for lead in leads:
        for c in lead.call_sessions.filter(provider__iexact='SARVAM').order_by('-created_at'):
            dur = round(c.duration_seconds, 1) if c.duration_seconds else 0.0
            loc_path = os.path.join(rec_cache_dir, f"{c.id}.wav")
            cached_text = f"Yes ({round(os.path.getsize(loc_path)/1024/1024, 2)} MB)" if os.path.exists(loc_path) else "No"

            matched = sarvam_data.get(c.provider_attempt_id)
            d_url = matched.get('audio_url') if matched else ""
            iid = matched.get('interaction_id') if matched else (c.metadata.get('interaction_id') if isinstance(c.metadata, dict) else "")
            if not d_url and iid and cfg.get('org_id'):
                enc_iid = str(iid).replace('/', '%2F')
                d_url = f"https://apps.sarvam.ai/api/analytics/v1/{cfg.get('org_id')}/{cfg.get('workspace_id')}/{cfg.get('app_id')}/recordings/{enc_iid}"

            b_url = f"{backend_public_url}/api/v1/voice-agent/recordings/{c.id}/audio/"

            v = c.final_agent_variables or {}
            c_summary = v.get('call_summary') or ""
            disp = v.get('disposition') or ""
            n_step = v.get('chosen_next_step') or ""
            n_dt = v.get('next_step_date_time') or ""

            t_lines = []
            if c.transcript and isinstance(c.transcript, list):
                for turn in c.transcript:
                    r = "Mahek (SWEAT AI)" if turn.get('role') == 'agent' else "Prospect"
                    t_lines.append(f"{r}: {turn.get('text') or turn.get('en_text') or ''}")
            transcript_str = "\n".join(t_lines)

            c_row = [
                f"{lead.first_name} {lead.last_name}".strip(),
                lead.phone_normalized,
                c.status,
                dur,
                d_url or "N/A",
                b_url,
                cached_text,
                c.provider_attempt_id or "",
                iid or "",
                fmt_dt(c.started_at or c.created_at),
                c_summary,
                disp,
                n_step,
                n_dt,
                transcript_str
            ]
            ws_calls.append(c_row)
            ws_calls.row_dimensions[call_row_idx].height = 24

            for col_i in range(1, len(c_row) + 1):
                cell = ws_calls.cell(row=call_row_idx, column=col_i)
                cell.font = font_data
                cell.border = thin_border
                cell.alignment = align_wrap if col_i in (11, 15) else align_left

            if d_url and d_url.startswith("http"):
                ws_calls.cell(row=call_row_idx, column=5).value = "⬇ Download Sarvam Audio"
                ws_calls.cell(row=call_row_idx, column=5).hyperlink = d_url
                ws_calls.cell(row=call_row_idx, column=5).font = font_link

            if b_url and b_url.startswith("http"):
                ws_calls.cell(row=call_row_idx, column=6).value = "▶ Stream Web Audio"
                ws_calls.cell(row=call_row_idx, column=6).hyperlink = b_url
                ws_calls.cell(row=call_row_idx, column=6).font = font_link

            call_row_idx += 1

    for col in ws_calls.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if '\n' in val:
                val = val.split('\n')[0]
            if len(val) > max_len:
                max_len = len(val)
        ws_calls.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 50)

    # Save to memory buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
